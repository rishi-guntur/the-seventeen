#!/usr/bin/env python3
"""
build_model.py — generates /model/basket-economics.xlsx for "The Seventeen".

Builds a basket-economics model for a rare-earth deposit:
  - Inputs   : deposit element distribution (% by oxide) + oxide prices
  - Calc     : weighted basket value per tonne + % contribution by element
  - Sensitivity : basket value under a price-floor scenario and an
                  export-control-shock scenario
  - Chart    : % contribution by element (so the memo can reference one chart)
  - Notes    : states that all numbers are illustrative until real figures
               are entered.

ALL element splits and prices in here are ILLUSTRATIVE PLACEHOLDERS, clearly
labelled, so the author can drop in real deposit data and live prices. The
spreadsheet is fully formula-driven: change the Inputs tab and Calc/Sensitivity
update automatically.

Usage:
    python scripts/build_model.py            # writes model/basket-economics.xlsx
    python scripts/build_model.py -o out.xlsx
"""

from __future__ import annotations

import argparse
import os
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Styling helpers
# ---------------------------------------------------------------------------
NAVY = "1F2A44"
SLATE = "33415C"
LIGHT = "EEF1F6"
PLACEHOLDER = "FFF3CD"      # soft amber = "you must fill this in"
ACCENT = "2E6DB4"
WHITE = "FFFFFF"

# Street conventions: hardcoded inputs = BLUE, formulas = BLACK,
# links to other tabs = GREEN.
INPUT_BLUE = "0000C0"
FORMULA_BLACK = "1A1A1A"
LINK_GREEN = "007B3A"

TITLE_FONT = Font(name="Calibri", size=14, bold=True, color=WHITE)
HEAD_FONT = Font(name="Calibri", size=11, bold=True, color=WHITE)
BODY_FONT = Font(name="Calibri", size=11, color=FORMULA_BLACK)
NOTE_FONT = Font(name="Calibri", size=10, italic=True, color="555555")
# Blue (Street convention for hardcodes) on the amber "fill me in" fill.
PLACEHOLDER_FONT = Font(name="Calibri", size=11, bold=True, color=INPUT_BLUE)
FORMULA_FONT = Font(name="Calibri", size=11, color=FORMULA_BLACK)
LINK_FONT = Font(name="Calibri", size=11, color=LINK_GREEN)

TITLE_FILL = PatternFill("solid", fgColor=NAVY)
HEAD_FILL = PatternFill("solid", fgColor=SLATE)
LIGHT_FILL = PatternFill("solid", fgColor=LIGHT)
PLACEHOLDER_FILL = PatternFill("solid", fgColor=PLACEHOLDER)

THIN = Side(style="thin", color="BFC7D5")
DOUBLE = Side(style="double", color=NAVY)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
# Output rows: single rule above, double rule below (Street convention).
TOTAL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=DOUBLE)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")

# ---------------------------------------------------------------------------
# Illustrative data (PLACEHOLDERS — author replaces these)
# ---------------------------------------------------------------------------
# The 17 elements, grouped. Splits below are ILLUSTRATIVE only.
# "Example Deposit A" ~ a light-dominant hard-rock basket (bastnasite-like).
# "Example Deposit B" ~ a basket with a richer heavy fraction (ionic-clay-like).
# These do NOT represent any real deposit. % by oxide should sum to ~100.
ELEMENTS = [
    # (Element, Oxide, Group, ExampleA_%, ExampleB_%, ExamplePrice_USD_per_kg)
    ("Lanthanum (La)",   "La2O3",  "Light", 28.0, 22.0, 1.0),
    ("Cerium (Ce)",      "CeO2",   "Light", 45.0, 40.0, 1.0),
    ("Praseodymium (Pr)","Pr6O11", "Light (magnet)", 4.5, 4.0, 90.0),
    ("Neodymium (Nd)",   "Nd2O3",  "Light (magnet)", 16.0, 14.0, 90.0),
    ("Samarium (Sm)",    "Sm2O3",  "Light", 2.0, 2.5, 3.0),
    ("Europium (Eu)",    "Eu2O3",  "Heavy", 0.3, 0.4, 28.0),
    ("Gadolinium (Gd)",  "Gd2O3",  "Heavy", 1.2, 2.0, 30.0),
    ("Terbium (Tb)",     "Tb4O7",  "Heavy (magnet)", 0.2, 1.5, 900.0),
    ("Dysprosium (Dy)",  "Dy2O3",  "Heavy (magnet)", 1.0, 6.0, 350.0),
    ("Holmium (Ho)",     "Ho2O3",  "Heavy", 0.15, 0.8, 90.0),
    ("Erbium (Er)",      "Er2O3",  "Heavy", 0.4, 2.5, 30.0),
    ("Thulium (Tm)",     "Tm2O3",  "Heavy", 0.05, 0.3, 100.0),
    ("Ytterbium (Yb)",   "Yb2O3",  "Heavy", 0.2, 1.5, 18.0),
    ("Lutetium (Lu)",    "Lu2O3",  "Heavy", 0.05, 0.3, 600.0),
    ("Yttrium (Y)",      "Y2O3",   "Heavy", 0.7, 1.5, 8.0),
    ("Scandium (Sc)",    "Sc2O3",  "Heavy", 0.0, 0.0, 1000.0),
    # Promethium (Pm) is radioactive / not naturally occurring at scale —
    # listed for completeness of "the 17" but assumed 0.
    ("Promethium (Pm)",  "Pm2O3",  "n/a",   0.0, 0.0, 0.0),
]

# Sensitivity scenario assumptions (ILLUSTRATIVE multipliers on the price column)
FLOOR_NDPR_MULT = 1.25      # price-floor scenario: NdPr supported above spot
SHOCK_HEAVY_MULT = 2.00     # export-control shock: heavies (Dy/Tb/etc.) spike
SHOCK_NDPR_MULT = 1.40      # NdPr also rises in a control shock


def _title_row(ws, text, ncols, subtitle=None):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=text)
    c.font = TITLE_FONT
    c.fill = TITLE_FILL
    c.alignment = LEFT
    ws.row_dimensions[1].height = 26
    for col in range(1, ncols + 1):
        ws.cell(row=1, column=col).fill = TITLE_FILL
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
        s = ws.cell(row=2, column=1, value=subtitle)
        s.font = NOTE_FONT
        s.alignment = LEFT
        ws.row_dimensions[2].height = 28


def build_inputs(ws):
    ncols = 7
    _title_row(
        ws, "INPUTS — Deposit element split (% by oxide) & oxide prices", ncols,
        subtitle=("ILLUSTRATIVE PLACEHOLDERS ONLY. Replace Example Deposit A/B "
                  "with [[INPUT: real % by oxide]] and the price column with "
                  "[[VERIFY: oxide prices as of <date>]]. Amber cells = you "
                  "must fill in. % columns should sum to ~100."),
    )
    headers = ["Element", "Oxide", "Group",
               "Example Deposit A\n(% by oxide) [[INPUT]]",
               "Example Deposit B\n(% by oxide) [[INPUT]]",
               "Oxide price\n(USD/kg) [[VERIFY: as of <date>]]",
               "Magnet element?"]
    hrow = 3
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=hrow, column=j, value=h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = CENTER
        c.border = BORDER
    ws.row_dimensions[hrow].height = 42

    r = hrow + 1
    for (elem, oxide, group, a, b, price) in ELEMENTS:
        ws.cell(row=r, column=1, value=elem).font = BODY_FONT
        ws.cell(row=r, column=2, value=oxide).font = BODY_FONT
        ws.cell(row=r, column=3, value=group).font = BODY_FONT
        # editable / placeholder cells
        for col, val in ((4, a / 100.0), (5, b / 100.0), (6, price)):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = PLACEHOLDER_FONT
            cell.fill = PLACEHOLDER_FILL
            cell.border = BORDER
            cell.alignment = RIGHT
            cell.number_format = "0.00%" if col in (4, 5) else '"$"#,##0.00'
        is_magnet = "magnet" in group.lower()
        ws.cell(row=r, column=7, value="Yes" if is_magnet else "").font = BODY_FONT
        for col in (1, 2, 3, 7):
            ws.cell(row=r, column=col).border = BORDER
            ws.cell(row=r, column=col).alignment = LEFT
        r += 1

    # totals row
    last = r - 1
    trow = r
    ws.cell(row=trow, column=3, value="TOTAL (should be ~100%)").font = Font(bold=True)
    for col in (4, 5):
        L = get_column_letter(col)
        t = ws.cell(row=trow, column=col,
                    value=f"=SUM({L}{hrow+1}:{L}{last})")
        t.font = Font(bold=True)
        t.number_format = "0.00%"
        t.fill = LIGHT_FILL
        t.border = BORDER
    ws.cell(row=trow, column=3).fill = LIGHT_FILL
    ws.cell(row=trow, column=3).border = BORDER

    # named-range helpers via column widths
    widths = [22, 10, 16, 16, 16, 18, 14]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    return {"hrow": hrow, "first": hrow + 1, "last": last, "ncols": ncols}


def build_calc(ws, ix):
    """Weighted basket value per tonne of contained oxide + % contribution."""
    first, last = ix["first"], ix["last"]
    ncols = 8
    _title_row(
        ws, "CALC — Weighted basket value per tonne & % contribution", ncols,
        subtitle=("Basket value per tonne of mixed rare-earth oxide (REO) = "
                  "Σ (oxide % × price USD/kg × 1,000 kg). Pulls live from the "
                  "Inputs tab. Deposit A is the default basket; switch the "
                  "formula column to Inputs col E to value Deposit B."),
    )
    headers = ["Element", "Group", "% by oxide (A)", "Price USD/kg",
               "kg per tonne REO", "Value USD/tonne", "% of basket value",
               "Magnet element?"]
    hrow = 3
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=hrow, column=j, value=h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = CENTER
        c.border = BORDER
    ws.row_dimensions[hrow].height = 30

    out_first = hrow + 1
    for i, src in enumerate(range(first, last + 1)):
        r = out_first + i
        ws.cell(row=r, column=1, value=f"=Inputs!A{src}").font = LINK_FONT
        ws.cell(row=r, column=2, value=f"=Inputs!C{src}").font = LINK_FONT
        pc = ws.cell(row=r, column=3, value=f"=Inputs!D{src}")
        pc.number_format = "0.00%"
        pc.font = LINK_FONT
        prc = ws.cell(row=r, column=4, value=f"=Inputs!F{src}")
        prc.number_format = '"$"#,##0.00'
        prc.font = LINK_FONT
        kg = ws.cell(row=r, column=5, value=f"=C{r}*1000")
        kg.number_format = "#,##0.0"
        kg.font = FORMULA_FONT
        val = ws.cell(row=r, column=6, value=f"=E{r}*D{r}")
        val.number_format = '"$"#,##0'
        val.font = FORMULA_FONT
        ws.cell(row=r, column=8, value=f"=Inputs!G{src}").font = LINK_FONT
        for col in range(1, 9):
            ws.cell(row=r, column=col).border = BORDER

    out_last = out_first + (last - first)
    total_row = out_last + 1
    # headline total
    tc = ws.cell(row=total_row, column=1, value="BASKET VALUE / TONNE REO  →")
    tc.font = Font(bold=True, size=12, color=NAVY)
    tv = ws.cell(row=total_row, column=6, value=f"=SUM(F{out_first}:F{out_last})")
    tv.font = Font(bold=True, size=12, color=NAVY)
    tv.number_format = '"$"#,##0'
    tv.fill = LIGHT_FILL
    tv.border = TOTAL_BORDER

    # % contribution column references the headline total
    for r in range(out_first, out_last + 1):
        pcell = ws.cell(row=r, column=7, value=f"=F{r}/$F${total_row}")
        pcell.number_format = "0.0%"
        pcell.font = FORMULA_FONT
        pcell.border = BORDER

    # magnet-element share callout
    mrow = total_row + 2
    ws.cell(row=mrow, column=1,
            value="Memo headline — magnet elements (NdPr+Dy+Tb) share of basket value:").font = Font(bold=True)
    # SUMIF over magnet rows
    mshare = ws.cell(
        row=mrow, column=6,
        value=(f'=SUMIF(H{out_first}:H{out_last},"Yes",F{out_first}:F{out_last})'
               f'/$F${total_row}'))
    mshare.number_format = "0.0%"
    mshare.font = Font(bold=True, color=ACCENT)

    widths = [22, 16, 14, 13, 15, 16, 16, 14]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    return {"hrow": hrow, "first": out_first, "last": out_last,
            "total_row": total_row, "ncols": ncols}


def build_sensitivity(ws, calc):
    first, last = calc["first"], calc["last"]
    ncols = 6
    _title_row(
        ws, "SENSITIVITY — Price-floor & export-control-shock scenarios", ncols,
        subtitle=(f"Illustrative scenario multipliers on the Inputs price column. "
                  f"Floor: NdPr ×{FLOOR_NDPR_MULT:g} (DoD-style price floor). "
                  f"Shock: heavies ×{SHOCK_HEAVY_MULT:g}, NdPr ×{SHOCK_NDPR_MULT:g} "
                  f"(China export-control squeeze). Replace multipliers with my "
                  f"own [[MY VIEW]] scenario assumptions."),
    )
    # assumption block
    arow = 3
    ws.cell(row=arow, column=1, value="Scenario assumptions (editable):").font = Font(bold=True)
    assumptions = [
        ("Floor — NdPr price multiplier", FLOOR_NDPR_MULT),
        ("Shock — heavy-RE price multiplier", SHOCK_HEAVY_MULT),
        ("Shock — NdPr price multiplier", SHOCK_NDPR_MULT),
    ]
    for k, (label, val) in enumerate(assumptions):
        r = arow + 1 + k
        ws.cell(row=r, column=1, value=label).font = BODY_FONT
        c = ws.cell(row=r, column=2, value=val)
        c.font = PLACEHOLDER_FONT
        c.fill = PLACEHOLDER_FILL
        c.border = BORDER
        c.number_format = "0.00x"
    floor_cell, heavy_cell, ndpr_cell = "B4", "B5", "B6"

    hrow = arow + len(assumptions) + 2
    headers = ["Element", "Group", "Base value\nUSD/tonne",
               "Floor scenario\nUSD/tonne", "Export-control shock\nUSD/tonne",
               "Magnet element?"]
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=hrow, column=j, value=h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = CENTER
        c.border = BORDER
    ws.row_dimensions[hrow].height = 36

    out_first = hrow + 1
    for i, src in enumerate(range(first, last + 1)):
        r = out_first + i
        ws.cell(row=r, column=1, value=f"=Calc!A{src}").font = LINK_FONT
        grp = f"Calc!B{src}"
        ws.cell(row=r, column=2, value=f"={grp}").font = LINK_FONT
        base = ws.cell(row=r, column=3, value=f"=Calc!F{src}")
        base.number_format = '"$"#,##0'
        base.font = LINK_FONT
        is_heavy = f'ISNUMBER(SEARCH("Heavy",Calc!B{src}))'
        # match the element symbol in parentheses, e.g. "(Nd)"/"(Pr)", so
        # "Promethium (Pm)" is not mis-caught by a bare "Pr" search.
        is_ndpr = (f'OR(ISNUMBER(SEARCH("(Nd)",Calc!A{src})),'
                   f'ISNUMBER(SEARCH("(Pr)",Calc!A{src})))')
        # Floor: NdPr lifted by floor multiplier, others unchanged
        floor = ws.cell(
            row=r, column=4,
            value=f"=Calc!F{src}*IF({is_ndpr},$ {floor_cell},1)".replace(" ", ""))
        floor.number_format = '"$"#,##0'
        floor.font = FORMULA_FONT
        # Shock: heavies × heavy mult, NdPr × ndpr mult, else base
        shock = ws.cell(
            row=r, column=5,
            value=(f"=Calc!F{src}*IF({is_ndpr},${ndpr_cell},"
                   f"IF({is_heavy},${heavy_cell},1))"))
        shock.number_format = '"$"#,##0'
        shock.font = FORMULA_FONT
        ws.cell(row=r, column=6, value=f"=Calc!H{src}").font = LINK_FONT
        for col in range(1, 7):
            ws.cell(row=r, column=col).border = BORDER

    out_last = out_first + (last - first)
    trow = out_last + 1
    ws.cell(row=trow, column=1, value="BASKET VALUE / TONNE  →").font = Font(bold=True, color=NAVY)
    for col, letter in ((3, "C"), (4, "D"), (5, "E")):
        t = ws.cell(row=trow, column=col,
                    value=f"=SUM({letter}{out_first}:{letter}{out_last})")
        t.font = Font(bold=True, size=12, color=NAVY)
        t.number_format = '"$"#,##0'
        t.fill = LIGHT_FILL
        t.border = TOTAL_BORDER

    # uplift line
    urow = trow + 1
    ws.cell(row=urow, column=1, value="Uplift vs base").font = Font(italic=True)
    fu = ws.cell(row=urow, column=4, value=f"=D{trow}/C{trow}-1")
    su = ws.cell(row=urow, column=5, value=f"=E{trow}/C{trow}-1")
    for c in (fu, su):
        c.number_format = "+0.0%;-0.0%"
        c.font = Font(italic=True, color=ACCENT)

    widths = [24, 16, 14, 16, 18, 14]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    return {"total_row": trow, "uplift_row": urow}


def build_summary(ws, calc, sens):
    """Cover / dashboard tab: title, TOC, key-outputs block, sensitivity chart."""
    ncols = 5
    _title_row(
        ws, "THE 17 — BASKET-ECONOMICS MODEL", ncols,
        subtitle=(f"Rare-earth basket value per tonne of mixed REO, with "
                  f"price-floor and export-control-shock scenarios. "
                  f"Prepared {date.today().isoformat()}. ALL FIGURES "
                  f"ILLUSTRATIVE until Inputs are replaced with real data — "
                  f"see Notes tab."),
    )

    # --- Table of contents ---
    r = 4
    ws.cell(row=r, column=1, value="TABLE OF CONTENTS").font = Font(
        name="Calibri", size=11, bold=True, color=NAVY)
    toc = [
        ("Summary", "This dashboard — key outputs & sensitivity chart."),
        ("Inputs", "Deposit % by oxide + oxide prices. BLUE cells = hardcoded "
                   "inputs to replace."),
        ("Calc", "Weighted basket value / tonne REO; % contribution by element."),
        ("Sensitivity", "Price-floor & export-control-shock scenarios."),
        ("Chart", "% contribution to basket value by element."),
        ("Notes", "Read me — conventions, caveats, fill-in checklist."),
    ]
    for i, (tab, desc) in enumerate(toc):
        rr = r + 1 + i
        t = ws.cell(row=rr, column=1, value=tab)
        t.font = Font(name="Calibri", size=11, bold=True, color=ACCENT)
        ws.cell(row=rr, column=2, value=desc).font = NOTE_FONT
        ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=ncols)

    # --- Key outputs block ---
    kr = r + len(toc) + 3
    ws.cell(row=kr, column=1, value="KEY OUTPUTS").font = Font(
        name="Calibri", size=11, bold=True, color=NAVY)
    ct = calc["total_row"]
    st = sens["total_row"]
    su = sens["uplift_row"]
    outputs = [
        ("Basket value / tonne REO (base)", f"=Calc!F{ct}", '"$"#,##0'),
        ("Basket value / tonne (price-floor scenario)", f"=Sensitivity!D{st}", '"$"#,##0'),
        ("Basket value / tonne (export-control shock)", f"=Sensitivity!E{st}", '"$"#,##0'),
        ("Uplift vs base — floor", f"=Sensitivity!D{su}", "+0.0%;-0.0%"),
        ("Uplift vs base — shock", f"=Sensitivity!E{su}", "+0.0%;-0.0%"),
        ("Magnet elements (NdPr+Dy+Tb) share of basket value",
         f"=Calc!F{ct + 2}", "0.0%"),
    ]
    for i, (label, formula, fmt) in enumerate(outputs):
        rr = kr + 1 + i
        lc = ws.cell(row=rr, column=1, value=label)
        lc.font = BODY_FONT
        lc.border = BORDER
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=3)
        for cc in range(1, 4):
            ws.cell(row=rr, column=cc).border = BORDER
        v = ws.cell(row=rr, column=4, value=formula)
        v.font = Font(name="Calibri", size=11, bold=True, color=LINK_GREEN)
        v.number_format = fmt
        v.alignment = RIGHT
        v.border = TOTAL_BORDER if i == 0 else BORDER
        v.fill = LIGHT_FILL
    ws.cell(row=kr + len(outputs) + 1, column=1,
            value="Green = pulled live from other tabs. Headline row "
                  "double-underlined. Figures illustrative until Inputs are real."
            ).font = NOTE_FONT

    # --- Sensitivity chart (native Excel) ---
    dr = kr + len(outputs) + 3
    ws.cell(row=dr, column=1, value="Scenario").font = HEAD_FONT
    ws.cell(row=dr, column=1).fill = HEAD_FILL
    ws.cell(row=dr, column=2, value="Basket value USD/tonne").font = HEAD_FONT
    ws.cell(row=dr, column=2).fill = HEAD_FILL
    scen = [("Base", f"=Calc!F{ct}"),
            ("Price floor", f"=Sensitivity!D{st}"),
            ("Export-control shock", f"=Sensitivity!E{st}")]
    for i, (label, formula) in enumerate(scen):
        rr = dr + 1 + i
        ws.cell(row=rr, column=1, value=label).font = BODY_FONT
        v = ws.cell(row=rr, column=2, value=formula)
        v.font = LINK_FONT
        v.number_format = '"$"#,##0'

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Basket value / tonne REO by scenario (illustrative)"
    chart.y_axis.title = "USD / tonne"
    chart.y_axis.numFmt = '"$"#,##0'
    data = Reference(ws, min_col=2, min_row=dr, max_row=dr + len(scen))
    cats = Reference(ws, min_col=1, min_row=dr + 1, max_row=dr + len(scen))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 9
    chart.width = 16
    chart.legend = None
    ws.add_chart(chart, f"F{kr}")

    widths = [40, 26, 8, 18, 4]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.sheet_view.showGridLines = False


def build_chart_sheet(ws, calc):
    ncols = 3
    _title_row(
        ws, "CHART — % contribution to basket value by element", ncols,
        subtitle="Auto-built from the Calc tab. The memo references this chart.",
    )
    # small data table the chart points at
    hrow = 3
    ws.cell(row=hrow, column=1, value="Element").font = HEAD_FONT
    ws.cell(row=hrow, column=2, value="% of basket value").font = HEAD_FONT
    ws.cell(row=hrow, column=1).fill = HEAD_FILL
    ws.cell(row=hrow, column=2).fill = HEAD_FILL
    first, last = calc["first"], calc["last"]
    n = last - first + 1
    for i in range(n):
        r = hrow + 1 + i
        ws.cell(row=r, column=1, value=f"=Calc!A{first+i}")
        pc = ws.cell(row=r, column=2, value=f"=Calc!G{first+i}")
        pc.number_format = "0.0%"
    data_last = hrow + n

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Contribution to basket value by element (illustrative)"
    chart.y_axis.title = "% of basket value"
    chart.x_axis.title = "Element"
    data = Reference(ws, min_col=2, min_row=hrow, max_row=data_last)
    cats = Reference(ws, min_col=1, min_row=hrow + 1, max_row=data_last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 12
    chart.width = 20
    chart.legend = None
    ws.add_chart(chart, "D3")

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16


def build_notes(ws):
    ncols = 1
    _title_row(ws, "NOTES — read me", ncols)
    notes = [
        "",
        "ALL NUMBERS IN THIS WORKBOOK ARE ILLUSTRATIVE PLACEHOLDERS.",
        "They do not represent any real deposit, price, or scenario and must "
        "not be circulated as analysis until replaced.",
        "",
        "To make this real, fill in:",
        "  1. Inputs!D (Example Deposit A) — [[INPUT]] real % by oxide for your "
        "deposit; rename the column to the deposit name.",
        "  2. Inputs!E (Example Deposit B) — [[INPUT]] a second real deposit, or "
        "delete the column.",
        "  3. Inputs!F (Oxide price) — [[VERIFY]] live oxide prices (USD/kg) "
        "with an as-of date; note China-domestic vs ex-China/CIF (memo §1.5).",
        "  4. Sensitivity!B4:B6 — [[MY VIEW]] your own floor / shock multipliers.",
        "",
        "How it flows: Inputs → Calc (weighted basket value/tonne + % "
        "contribution) → Sensitivity (floor & export-control scenarios) → Chart.",
        "Everything is formula-linked: change Inputs and the rest updates.",
        "",
        "Color conventions (Street standard):",
        "  BLUE text = hardcoded input (yours to change; amber fill = must fill in).",
        "  BLACK text = formula within the tab. Do not overwrite.",
        "  GREEN text = link pulling from another tab. Do not overwrite.",
        "  Double-underlined rows = final outputs (basket value / tonne).",
        "",
        "Headline number to quote in the memo: Calc!F<total row> "
        "(BASKET VALUE / TONNE REO). It is illustrative until step 1–3 are done.",
        "",
        "Pricing caveat (mirror of memo §1.5 / Sources): rare-earth pricing is "
        "opaque, thin, two-regime, and goes stale fast. Date every price.",
        "",
        f"Workbook generated by scripts/build_model.py on {date.today().isoformat()}.",
    ]
    r = 3
    for line in notes:
        c = ws.cell(row=r, column=1, value=line)
        if line.isupper() and line:
            c.font = Font(bold=True, color="8A2A2A")
        elif line.startswith("  "):
            c.font = PLACEHOLDER_FONT
        else:
            c.font = NOTE_FONT
        c.alignment = LEFT
        r += 1
    ws.column_dimensions["A"].width = 100


def main():
    parser = argparse.ArgumentParser(description="Build basket-economics.xlsx")
    parser.add_argument("-o", "--output",
                        default=os.path.join("model", "basket-economics.xlsx"))
    args = parser.parse_args()

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Summary"
    ws_inputs = wb.create_sheet("Inputs")
    ws_calc = wb.create_sheet("Calc")
    ws_sens = wb.create_sheet("Sensitivity")
    ws_chart = wb.create_sheet("Chart")
    ws_notes = wb.create_sheet("Notes")

    ix = build_inputs(ws_inputs)
    calc = build_calc(ws_calc, ix)
    sens = build_sensitivity(ws_sens, calc)
    build_summary(ws_summary, calc, sens)
    build_chart_sheet(ws_chart, calc)
    build_notes(ws_notes)

    # freeze header rows
    for ws, cell in ((ws_inputs, "A4"), (ws_calc, "A4"), (ws_sens, "A4")):
        ws.freeze_panes = cell

    os.makedirs(os.path.dirname(os.path.abspath(args.output)), exist_ok=True)
    wb.save(args.output)
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
